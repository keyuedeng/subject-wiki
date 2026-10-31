"""Read just enough of a file to classify it: the first page(s), slide(s) or lines.

This is not parsing (phase 2); it never fails an upload. Unreadable files give an
empty Peek with `error` set, and the classifier falls back to the filename.
"""

import io
import logging
from dataclasses import dataclass

logger = logging.getLogger(__name__)

MAX_CHARS = 8000
PDF_PAGES = 2
PPTX_SLIDES = 2
DOCX_PARAGRAPHS = 30
TEXT_LINES = 200
# Title candidates come from a little further in: page 1 is often a copyright notice.
TITLE_PAGES = 3
TITLE_CANDIDATES_PER_PAGE = 4
TITLE_SIZE_RATIO = 1.15  # a title is set at least this much larger than body text


@dataclass(frozen=True)
class Peek:
    text: str = ""
    lines: tuple[str, ...] = ()  # only for plain-text formats
    landscape: bool | None = None  # first page wider than tall (PDF/PPTX)
    # Likely titles in preference order (largest text first); classify.titles picks one.
    title_candidates: tuple[str, ...] = ()
    error: str | None = None


def peek(data: bytes, ext: str) -> Peek:
    try:
        match ext:
            case "pdf":
                return _peek_pdf(data)
            case "pptx":
                return _peek_pptx(data)
            case "docx":
                return _peek_docx(data)
            case "ipynb":
                return _peek_ipynb(data)
            case "txt" | "vtt" | "srt":
                return _peek_text(data)
            case _:
                return Peek()
    except Exception as exc:  # corrupt or password-protected files, unusual encodings
        logger.info("peek failed for .%s file: %s", ext, type(exc).__name__)
        return Peek(error=f"could not read .{ext} contents ({type(exc).__name__})")


def _peek_pdf(data: bytes) -> Peek:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.page_count == 0:
            return Peek(error="PDF has no pages")
        first = doc[0].rect
        text = "\n".join(doc[i].get_text() for i in range(min(PDF_PAGES, doc.page_count)))
        candidates: list[str] = []
        for i in range(min(TITLE_PAGES, doc.page_count)):
            candidates += _pdf_page_titles(doc[i])
        return Peek(
            text=text[:MAX_CHARS],
            landscape=first.width > first.height,
            title_candidates=tuple(candidates),
        )


def _pdf_page_titles(page) -> list[str]:
    """Lines set larger than the page's body text, largest first. Adjacent same-size
    lines are joined so a title that wraps onto two lines stays whole."""
    lines: list[tuple[float, str]] = []
    chars_by_size: dict[float, int] = {}
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            spans = [s for s in line["spans"] if s["text"].strip()]
            if not spans:
                continue
            size = round(max(s["size"] for s in spans), 1)
            text = " ".join(s["text"].strip() for s in spans)
            chars_by_size[size] = chars_by_size.get(size, 0) + len(text)
            if lines and lines[-1][0] == size:
                lines[-1] = (size, f"{lines[-1][1]} {text}")
            else:
                lines.append((size, text))
    if not lines:
        return []
    # Body text is the size with the most characters, unless that is the largest size
    # (a sparse title slide), then the next size down. One size throughout: no title.
    sizes = sorted(chars_by_size)
    body = max(chars_by_size, key=lambda size: chars_by_size[size])
    if body == sizes[-1] and len(sizes) > 1:
        body = sizes[-2]
    titles = [(size, text) for size, text in lines if size >= body * TITLE_SIZE_RATIO]
    titles.sort(key=lambda item: -item[0])  # stable: equal sizes keep page order
    return [text for _, text in titles[:TITLE_CANDIDATES_PER_PAGE]]


def _peek_pptx(data: bytes) -> Peek:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = list(prs.slides)
    chunks = []
    for slide in slides[:PPTX_SLIDES]:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    titles = []
    for slide in slides[:TITLE_PAGES]:
        title = slide.shapes.title
        if title is not None and title.has_text_frame and title.text_frame.text.strip():
            titles.append(" ".join(title.text_frame.text.split()))
    landscape = None
    if prs.slide_width and prs.slide_height:
        landscape = prs.slide_width > prs.slide_height
    return Peek(
        text="\n".join(chunks)[:MAX_CHARS], landscape=landscape, title_candidates=tuple(titles)
    )


def _peek_docx(data: bytes) -> Peek:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paragraphs = [p for p in doc.paragraphs[:DOCX_PARAGRAPHS] if p.text.strip()]
    headings = [
        p.text.strip()
        for p in paragraphs
        if (p.style is not None and p.style.name.startswith(("Title", "Heading")))
    ]
    short = [p.text.strip() for p in paragraphs[:5] if len(p.text.strip()) <= 90]
    return Peek(
        text="\n".join(p.text for p in paragraphs)[:MAX_CHARS],
        title_candidates=tuple(headings[:3] + short[:2]),
    )


def _peek_ipynb(data: bytes) -> Peek:
    """Jupyter notebook: markdown and code cells' text; markdown headings are titles."""
    import json

    cells = json.loads(data.decode("utf-8-sig")).get("cells") or []
    texts, headings = [], []
    for cell in cells:
        source = cell.get("source") or ""
        source = "".join(source) if isinstance(source, list) else str(source)
        texts.append(source)
        if cell.get("cell_type") == "markdown":
            headings += [ln.lstrip("#").strip() for ln in source.splitlines()
                         if ln.startswith("#") and ln.lstrip("#").strip()]  # fmt: skip
    return Peek(text="\n".join(texts)[:MAX_CHARS], title_candidates=tuple(headings[:3]))


def _peek_text(data: bytes) -> Peek:
    head = data[: MAX_CHARS * 4].decode("utf-8-sig", errors="replace")
    lines = tuple(head.splitlines()[:TEXT_LINES])
    return Peek(text=head[:MAX_CHARS], lines=lines)
