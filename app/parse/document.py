"""Open a slide deck, render its pages and measure what is on each one.

Measurements ignore the deck's *template*: images, vector shapes and text lines that
repeat on most pages (logos, header bars, running titles, footers). Without that, a
logo on every slide would make every page look like it has an image.
"""

import io
import logging
import re
import shutil
import subprocess
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf

logger = logging.getLogger(__name__)

RENDER_DPI = 150
TEMPLATE_SHARE = 0.6  # on at least this share of pages -> template
MIN_PAGES_FOR_TEMPLATE = 4  # too few pages to tell template from content
MONO_FONT_RE = re.compile(r"mono|courier|consolas|menlo|inconsolata|code|fixed", re.IGNORECASE)
MATH_FONT_RE = re.compile(r"cmmi|cmsy|cmex|msam|msbm|cambria ?math|stix|symbol", re.IGNORECASE)
PAGE_NUMBER_RE = re.compile(r"^\s*\d{1,3}\s*$")
SOFFICE_PATHS = ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice")
# Hidden slides are left out explicitly, so PDF pages line up with pptx_notes().
PDF_EXPORT = 'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"false"}}'


class ConversionError(Exception):
    pass


@dataclass(frozen=True)
class Line:
    text: str
    size: float
    x0: float
    y0: float
    mono: bool
    bold: bool


@dataclass
class PageInfo:
    page_no: int  # 1-based
    width: float
    height: float
    text_raw: str  # the full text layer, template included
    lines: list[Line]  # non-template lines, top to bottom
    image_area: float  # share of the page covered by non-template images (0-1)
    image_count: int
    drawing_count: int  # non-template vector paths
    drawing_area: float  # share of the page covered by their bounding box (0-1)
    drawing_sigs: frozenset = field(default_factory=frozenset)
    image_sigs: frozenset = field(default_factory=frozenset)
    math_font: bool = False

    @property
    def body_lines(self) -> list[str]:
        return [ln.text for ln in self.lines]

    @property
    def body_chars(self) -> int:
        return sum(len(t) for t in self.body_lines)

    @property
    def title_lines(self) -> list[Line]:
        """The largest text on the page (at most two lines, for a wrapped title)."""
        if not self.lines:
            return []
        biggest = max(ln.size for ln in self.lines)
        return [ln for ln in self.lines if ln.size >= biggest - 0.5][:2]

    @property
    def title(self) -> str | None:
        return " ".join(ln.text for ln in self.title_lines) or None

    @property
    def signals(self) -> dict:
        return {
            "body_chars": self.body_chars,
            "image_count": self.image_count,
            "image_area": round(self.image_area, 3),
            "drawing_count": self.drawing_count,
            "drawing_area": round(self.drawing_area, 3),
            "math_font": self.math_font,
        }


@dataclass
class Deck:
    pdf: bytes
    pages: list[PageInfo]
    notes: dict[int, str]  # page_no -> speaker notes (PPTX only)


def deck_from_pdf(pdf: bytes, notes: dict[int, str] | None = None,
                  expected_pages: int | None = None) -> Deck:  # fmt: skip
    """Measure a PDF. Speaker notes are matched to pages by position, so they are only
    attached when the PDF has exactly one page per visible slide."""
    pages = measure(pdf)
    if notes and expected_pages is not None and expected_pages != len(pages):
        logger.warning(
            "PPTX has %d visible slides but the PDF has %d pages; speaker notes skipped",
            expected_pages,
            len(pages),
        )
        notes = {}
    return Deck(pdf=pdf, pages=pages, notes=notes or {})


def render_page(doc: pymupdf.Document, page_no: int) -> bytes:
    return doc[page_no - 1].get_pixmap(dpi=RENDER_DPI).tobytes("png")


# PPTX


def find_soffice() -> str | None:
    for candidate in SOFFICE_PATHS:
        if found := shutil.which(candidate) or (Path(candidate).is_file() and candidate):
            return found
    return None


def pptx_to_pdf(data: bytes) -> bytes:
    soffice = find_soffice()
    if soffice is None:
        raise ConversionError(
            "PPTX needs LibreOffice to convert (brew install --cask libreoffice), "
            "or upload the PDF export instead."
        )
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "deck.pptx"
        src.write_bytes(data)
        # A private profile dir avoids clashing with a running LibreOffice instance.
        profile = f"-env:UserInstallation=file://{tmp}/profile"
        result = subprocess.run(
            [soffice, profile, "--headless", "--convert-to", PDF_EXPORT, "--outdir", tmp, str(src)],
            capture_output=True,
            timeout=300,
        )
        out = Path(tmp) / "deck.pdf"
        if result.returncode != 0 or not out.exists():
            raise ConversionError("LibreOffice could not convert the PPTX to PDF")
        return out.read_bytes()


def pptx_notes(data: bytes) -> tuple[dict[int, str], int]:
    """Speaker notes by page number, and the number of visible slides. Hidden slides
    are skipped, as in the PDF export."""
    from pptx import Presentation

    notes: dict[int, str] = {}
    page_no = 0
    for slide in Presentation(io.BytesIO(data)).slides:
        if slide._element.get("show") == "0":
            continue
        page_no += 1
        if slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text.strip()
            if text:
                notes[page_no] = text
    return notes, page_no


# Measuring


def _round_rect(r: pymupdf.Rect) -> tuple:
    return tuple(round(v / 4) for v in (r.x0, r.y0, r.x1, r.y1))  # 4pt grid


def _drawing_sig(d: dict) -> tuple:
    return (d.get("type"), _round_rect(d["rect"]), d.get("fill") is not None)


def _image_sig(xref: int, rect: pymupdf.Rect) -> tuple:
    return (xref, _round_rect(rect))


def _norm(text: str) -> str:
    return " ".join(text.split()).lower()


@dataclass
class _Raw:
    page: pymupdf.Page
    lines: list[Line]
    images: list[tuple[tuple, pymupdf.Rect]]
    drawings: list[tuple[tuple, pymupdf.Rect]]
    math_font: bool


BULLET_ONLY_RE = re.compile(r"^[•●▪■◦○‣∙·\-–—*]$")


def _is_maths(span: dict) -> bool:
    """A maths font drawing real symbols. LibreOffice draws PowerPoint bullets and arrows
    with OpenSymbol private-use glyphs (U+E000-U+F8FF); those aren't equations."""
    if not MATH_FONT_RE.search(span["font"]):
        return False
    return any(not ("\ue000" <= ch <= "\uf8ff") and not ch.isspace() and ch not in "•●▪■◦○‣∙·"
               for ch in span["text"])  # fmt: skip


def _raw_page(page: pymupdf.Page) -> _Raw:
    """Lines in reading order: text blocks top-to-bottom then left-to-right (so columns
    stay together), lines in order within each block."""
    lines: list[Line] = []
    math = False
    blocks = [b for b in page.get_text("dict")["blocks"] if b.get("lines")]
    blocks.sort(key=lambda b: (round(b["bbox"][1] / 10), b["bbox"][0]))
    for block in blocks:
        pending_bullet: str | None = None
        for line in block["lines"]:
            spans = [s for s in line["spans"] if s["text"].strip()]
            if not spans:
                continue
            # Span texts carry their own spacing; joining with "" keeps "draw()" intact.
            text = " ".join("".join(s["text"] for s in line["spans"]).split())
            if BULLET_ONLY_RE.match(text):
                pending_bullet = text  # a bullet exported as its own line
                continue
            if pending_bullet:
                text, pending_bullet = f"{pending_bullet} {text}", None
            fonts = [s["font"] for s in spans]
            math = math or any(_is_maths(s) for s in spans)
            lines.append(
                Line(
                    text=text,
                    size=round(max(s["size"] for s in spans), 1),
                    x0=line["bbox"][0],
                    y0=line["bbox"][1],
                    mono=all(MONO_FONT_RE.search(f) for f in fonts),
                    bold=all(s["flags"] & 16 or "bold" in s["font"].lower() for s in spans),
                )
            )
    images = []
    for img in page.get_images(full=True):
        for rect in page.get_image_rects(img[0]):
            images.append((_image_sig(img[0], rect), rect))
    drawings = [(_drawing_sig(d), d["rect"]) for d in page.get_drawings()]
    return _Raw(page, lines, images, drawings, math)


def measure(pdf: bytes) -> list[PageInfo]:
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        raws = [_raw_page(page) for page in doc]
        template_lines, template_images, template_drawings = _template(raws)
        pages = []
        for i, raw in enumerate(raws):
            rect = raw.page.rect
            area = rect.width * rect.height or 1
            lines = [
                ln
                for ln in raw.lines
                if _norm(ln.text) not in template_lines and not PAGE_NUMBER_RE.match(ln.text)
            ]
            imgs = [(s, r) for s, r in raw.images if s not in template_images and s[0] not in
                    template_images]  # fmt: skip
            draws = [(s, r) for s, r in raw.drawings if s not in template_drawings]
            image_area = min(1.0, sum((r & rect).get_area() for _, r in imgs) / area)
            drawing_area = 0.0
            if draws:
                box = pymupdf.Rect(draws[0][1])
                for _, r in draws[1:]:
                    box |= r
                drawing_area = min(1.0, (box & rect).get_area() / area)
            pages.append(
                PageInfo(
                    page_no=i + 1,
                    width=rect.width,
                    height=rect.height,
                    text_raw=raw.page.get_text(),
                    lines=lines,
                    image_area=image_area,
                    image_count=len(imgs),
                    drawing_count=len(draws),
                    drawing_area=drawing_area,
                    drawing_sigs=frozenset(s for s, _ in draws),
                    image_sigs=frozenset(s for s, _ in imgs),
                    math_font=raw.math_font,
                )
            )
        return pages


def _template(raws: list[_Raw]) -> tuple[set, set, set]:
    """Elements repeated on at least TEMPLATE_SHARE of pages."""
    n = len(raws)
    if n < MIN_PAGES_FOR_TEMPLATE:
        return set(), set(), set()
    threshold = max(2, TEMPLATE_SHARE * n)

    def frequent(per_page: list[set]) -> set:
        counts = Counter(x for items in per_page for x in items)
        return {x for x, c in counts.items() if c >= threshold}

    lines = frequent([{_norm(ln.text) for ln in r.lines} for r in raws])
    # An image counts as template by placement or by xref alone (same logo, moved).
    images = frequent([{s for s, _ in r.images} for r in raws])
    images |= frequent([{s[0] for s, _ in r.images} for r in raws])
    drawings = frequent([{s for s, _ in r.drawings} for r in raws])
    return lines, images, drawings
