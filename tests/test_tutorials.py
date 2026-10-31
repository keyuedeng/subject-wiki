"""Tutorials: reading sheets (PDF, DOCX, .ipynb), splitting them into sections, labelling
(type, concepts, lecture coverage), solutions, and the board card and viewer. The model
and the lecture search are faked."""

import io
import json
import threading
import uuid

import pymupdf
import pytest
from sqlalchemy import select

from app.db.models import TutorialSection
from app.llm.client import StructuredResponse, Usage
from app.parse.document import deck_from_pdf, find_soffice
from app.parse.pipeline import Parser
from app.qa.labels import SourceRef
from app.services.uploads import ingest_upload
from app.tutorials.blocks import ipynb_blocks, pdf_blocks
from app.tutorials.label import practised_weeks
from app.tutorials.pipeline import TutorialProcessor, TutorialSettings
from app.tutorials.sections import split_sections
from tests.search_fixtures import make_llm
from tests.test_pipeline import FakeVision

# A sheet laid out like SWEN30006 Workshop 4: title, setup, background, exercises with
# sub-parts, then a short part heading right before an exercise.
SHEET = [
    (18, False, "Workshop 4: Fundamental Design Patterns"),
    (14, True, "Requisite Knowledge and Tools"),
    (10, False, "You should be familiar with UML class diagrams and the GRASP patterns "
                "before this workshop. Install a UML editor before you come to class."),
    (14, True, "Introduction"),
    (10, False, "GRASP is a set of patterns for assigning responsibilities to classes. "
                "Our focus today is a simple online marketplace where sellers list items."),
    (14, True, "Exercise One - Determining Experts"),
    (10, False, "Use the Information Expert pattern to decide who should do the following:"),
    (10, False, "(a) Determining the highest bidder in an auction at any given time."),
    (10, False, "(b) Finding all the items within a subcategory of the marketplace."),
    (14, True, "Part B"),
    (14, True, "Exercise Two - Gossip"),
    (10, False, "Use a gossip protocol to spread listing updates between the marketplace "
                "servers, and explain how many rounds it takes to reach every server."),
]  # fmt: skip
ANSWERS = [
    (18, False, "Workshop 4: Partial Solution"),
    (14, True, "Exercise One - Determining Experts"),
    (10, False, "The Listing is the expert for the highest bidder: it has all the bids. "
                "SubCategory has the most direct access to the items within it."),
    (14, True, "Exercise Two - Gossip"),
    (10, False, "Each round doubles the servers that know, so it takes about log2(n) rounds."),
]  # fmt: skip


def make_sheet(lines) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    y = 60
    for size, bold, text in lines:
        font = "hebo" if bold else "helv"
        width = 480
        rect = pymupdf.Rect(60, y, 60 + width, y + 400)
        used = page.insert_textbox(rect, text, fontsize=size, fontname=font)
        y += (400 - used) + size * 0.8
    return doc.tobytes()


# Splitting


def test_pdf_sheet_splits_by_its_headings_and_keeps_sub_parts():
    sections = split_sections(pdf_blocks(deck_from_pdf(make_sheet(SHEET)), {}))
    # The title block has no text of its own: it opens the first section.
    assert [s.label for s in sections] == ["Requisite Knowledge and Tools", "Introduction",
                                           "Q1", "Q2"]  # fmt: skip
    assert "Workshop 4: Fundamental Design Patterns" in sections[0].markdown()
    q1, q2 = sections[2], sections[3]
    assert (q1.number, q1.is_exercise) == (1, True)
    md = q1.markdown()
    assert "(a) Determining the highest bidder" in md and "(b) Finding all the items" in md
    # "Part B" had no text of its own: it joined the exercise after it.
    assert q2.number == 2 and "Part B" in q2.markdown() and "Part B" not in md


def test_notebook_headings_markdown_and_code():
    nb = {"cells": [
        {"cell_type": "markdown", "source": [
            "# Lab 3: Hashing\n", "A salt is a random value stored with each password hash, "
            "so two users with the same password get different hashes."]},
        {"cell_type": "code", "source": ["import hashlib\n", "print(1)"],
         "outputs": [{"text": "1"}]},
        {"cell_type": "markdown", "source": ["**Question 1** Crack the hash below."]},
        {"cell_type": "markdown", "source": "## Question 2\nUse a mask attack on the PINs."},
    ], "nbformat": 4}  # fmt: skip
    blocks = ipynb_blocks(json.dumps(nb).encode())
    assert [b.kind for b in blocks] == ["heading", "text", "code", "heading", "text",
                                         "heading", "text"]  # fmt: skip
    sections = split_sections(blocks)
    assert [s.label for s in sections] == ["Lab 3: Hashing", "Q1", "Q2"]
    assert "import hashlib" in sections[0].markdown() and '"text": "1"' not in str(blocks)


@pytest.mark.skipif(find_soffice() is None, reason="needs LibreOffice")
def test_docx_sheet_splits_by_its_heading_styles():
    from docx import Document

    from app.parse.document import docx_to_pdf

    doc = Document()
    doc.add_heading("Tutorial 2: Threat Analysis", level=0)
    doc.add_heading("Background", level=1)
    doc.add_paragraph("An attack tree breaks an attacker's goal into the ways it can be "
                      "reached, so defenders can see which branches are cheapest.")  # fmt: skip
    doc.add_heading("Task 1", level=1)
    doc.add_paragraph("Draw an attack tree for stealing a student's exam answers.")
    doc.add_paragraph("Mark the cheapest branch and say why it is the cheapest one.")
    doc.add_heading("Task 2", level=1)
    doc.add_paragraph("Use STRIDE to list threats to the subject's online quiz system.")
    buf = io.BytesIO()
    doc.save(buf)
    deck = deck_from_pdf(docx_to_pdf(buf.getvalue()), exact_spaces=True)
    labels = [s.label for s in split_sections(pdf_blocks(deck, {}))]
    assert labels == ["Background", "Q1", "Q2"]  # a one-line exercise is still its own


# Labelling and solutions

LECTURES = uuid.uuid4()


def ref(week, page, title, body):
    meta = {"page_no": page, "slide_id": str(uuid.uuid4())}
    return SourceRef(uuid.uuid5(LECTURES, f"{week}:{page}"), "slide", LECTURES, "GRASP", week,
                     week, page, title, body, meta)  # fmt: skip


FOUND = {
    "Information Expert": [
        ref(
            4,
            19,
            "Information Expert",
            "Assign a responsibility to the class with the information.",
        )
    ],
    "Creator": [ref(4, 13, "Creator", "B creates A when B contains or aggregates A.")],
    "gossip protocol": [],  # not lectured, and the sheet explains it
    "bidding": [],  # not lectured, and the sheet only uses it
    "Adapter": [ref(7, 6, "Adapter", "Adapter converts one interface into another.")],
}


class FakeModel:
    name = "fake"

    def __init__(self):
        self.calls: list[str] = []
        self._lock = threading.Lock()

    def structured(self, *, model, instructions, content, schema, effort, max_output_tokens,
                   cache_hint=None):  # fmt: skip
        text = content[-1]["text"]
        with self._lock:
            self.calls.append(schema.__name__)
        if schema.__name__ == "TutorialSections":
            items = []
            for sid, heading in _sections_in(text):
                if "Requisite" in heading or heading.startswith("Workshop"):
                    items.append({"id": sid, "type": "setup_admin", "concepts": ["UML"]})
                elif heading.startswith("Exercise One"):
                    items.append(
                        {
                            "id": sid,
                            "type": "exercise",
                            "concepts": ["Information Expert", "bidding"],
                        }
                    )
                elif heading.startswith("Exercise Two"):
                    items.append(
                        {"id": sid, "type": "exercise", "concepts": ["gossip protocol", "Adapter"]}
                    )
                else:
                    items.append({"id": sid, "type": "explanation", "concepts": ["Creator"]})
            response = {"sections": items}
        else:
            concepts = []
            for cid, name, eids in _concepts_in(text):
                if eids:
                    concepts.append({"id": cid, "status": "covered", "evidence": eids[:1],
                                     "summary": ""})  # fmt: skip
                else:
                    summary = (
                        "Servers spread updates by telling random peers each round."
                        if name == "gossip protocol"
                        else ""
                    )
                    concepts.append({"id": cid, "status": "not_covered", "evidence": [],
                                     "summary": summary})  # fmt: skip
            response = {"concepts": concepts}
        return StructuredResponse(schema.model_validate(response), Usage(3000, 0, 400, 0))

    def embed(self, *, model, texts, dimensions):
        return [[1.0] * 8 for _ in texts], 0


def _sections_in(text):
    import re

    return re.findall(r"^\[(T\d+)\] (.*)$", text, re.MULTILINE)


def _concepts_in(text):
    import re

    out = []
    for block in text.split("\n\n[C")[1:]:
        head, *rest = block.split("\n")
        cid, name = re.match(r"(\d+)\] (.*)", head).groups()
        eids = re.findall(r"^\s+\[(E\d+)\]", "\n".join(rest), re.MULTILINE)
        out.append((f"C{cid}", name, eids))
    return out


@pytest.fixture
def model():
    return FakeModel()


@pytest.fixture
def processor(session_factory, storage, model):
    return TutorialProcessor(
        session_factory, storage, make_llm(session_factory, model),
        Parser(session_factory, storage, FakeVision),
        search=lambda db, subject, query: FOUND.get(query, []),
        settings=TutorialSettings("gpt-6.1-sol", "low", concurrency=2),
    )  # fmt: skip


@pytest.fixture
def sheets(db, storage, subject):
    """Workshop 4 filed in week 5 (the week it's taught), with its solutions."""

    def up(name, data, slot):
        return ingest_upload(db, storage, subject, name, None, data, max_bytes=50_000_000,
                             slot=slot, slot_week=5).file  # fmt: skip

    return (up("Workshop_04.pdf", make_sheet(SHEET), "tutorials"),
            up("Partial_Solution_04.pdf", make_sheet(ANSWERS), "tutorials"))  # fmt: skip


def rows(db, f):
    db.expire_all()
    q = select(TutorialSection).where(TutorialSection.file_id == f.id)
    return {s.label: s for s in db.scalars(q.order_by(TutorialSection.idx))}


def test_tutorial_is_split_labelled_and_linked_to_lectures_and_solutions(db, processor, model,
                                                                         sheets):  # fmt: skip
    tute, sols = sheets
    assert (tute.type, tute.week, sols.type) == ("tutorial", 5, "solutions")
    processor.process(sols.id)
    processor.process(tute.id)
    db.refresh(tute)
    assert tute.status == "parsed", tute.status_detail
    assert "4 sections (2 exercises) · 4 labelled" in tute.status_detail
    s = rows(db, tute)

    setup = s["Requisite Knowledge and Tools"]
    assert (setup.section_type, setup.concepts, setup.coverage) == ("setup_admin", [], None)

    q1 = s["Q1"]
    assert (q1.section_type, q1.coverage, q1.earliest_week) == ("exercise", "new_concept", 4)
    expert, bidding = q1.concepts
    assert expert["status"] == "covered" and expert["evidence"][0]["label"] == "Week 4 · slide 19"
    assert expert["evidence"][0]["href"].startswith("/slides/")
    assert (bidding["status"], bidding["summary"]) == ("not_covered", "")  # used, not explained

    q2 = s["Q2"]
    gossip, adapter = q2.concepts
    assert gossip["summary"].startswith("Servers spread updates")  # from the sheet only
    assert adapter["earliest_week"] == 7  # lectured after week 5: ahead of lectures

    assert s["Introduction"].coverage == "application"
    # Weeks it practises come from the evidence, not from the week it's filed in; one
    # concept lectured in week 7 doesn't make week 7 a week it practises.
    assert practised_weeks(list(s.values())) == [4]

    answers = rows(db, sols)
    assert (q1.solution_source, q1.solution_section_id) == ("official", answers["Q1"].id)
    assert q2.solution_section_id == answers["Q2"].id

    # Processing again makes no model calls.
    calls = len(model.calls)
    processor.process(tute.id)
    assert len(model.calls) == calls
    db.refresh(tute)
    assert "this run $0.00" in tute.status_detail


def test_solutions_link_when_uploaded_later(db, storage, subject, processor, sheets):
    tute, sols = sheets
    processor.process(tute.id)
    assert rows(db, tute)["Q1"].solution_source == "none"
    processor.process(sols.id)
    assert rows(db, tute)["Q1"].solution_source == "official"


def test_board_card_and_viewer(db, processor, sheets, client, subject):
    tute, sols = sheets
    processor.process(sols.id)
    processor.process(tute.id)
    board = client.get(f"/subjects/{subject.id}/view/board").text
    week5 = board.split(">Week 5<", 1)[1].split(">Week 6<", 1)[0]
    assert "Practises Week 4 lectures" in week5
    assert f'href="/files/{tute.id}/tutorial"' in week5
    assert "4 sections (2 exercises) · labelled" in week5 and "2 exercises matched" in week5

    page = client.get(f"/files/{tute.id}/tutorial").text
    assert "filed in Week 5" in page and "Practises Week 4 lectures" in page
    assert 'class="cite cite-slide" href="/slides/' in page and "Week 4 · slide 19" in page
    assert "Ahead of lectures: taught in Week 7" in page
    assert "Used in tutorial, not explained in your material." in page
    assert "Servers spread updates by telling random peers" in page
    assert "Show solution" in page and "The Listing is the expert" in page
    assert "1 with a new concept" not in page and "2 with a new concept" in page
    sol_page = client.get(f"/files/{sols.id}/tutorial").text
    assert "Answers <a" in sol_page and "Workshop_04.pdf" in sol_page


def test_practised_weeks_need_two_concepts_and_mostly_new_sheets_say_so():
    from app.tutorials.label import mostly_new

    def c(name, *weeks):
        ev = [{"week": w} for w in weeks]
        return {"name": name, "status": "covered" if ev else "not_covered", "evidence": ev}

    stray = [{"concepts": [c("Security controls", 2), c("Spectre"), c("Rowhammer"),
                           c("Heartbleed"), c("Dead stores")]}]  # fmt: skip
    assert practised_weeks(stray) == [] and mostly_new(stray)
    grasp = [{"concepts": [c("Creator", 4), c("Information Expert", 4), c("creator", 4, 7),
                           c("Controller")]}]  # fmt: skip
    assert practised_weeks(grasp) == [4] and not mostly_new(grasp)


def test_a_later_deck_showing_two_diagrams_isnt_a_practised_week():
    def c(name, *weeks):
        return {"name": name, "status": "covered", "evidence": [{"week": w} for w in weeks]}

    sheet = [{"concepts": [c("Creator", 4), c("Information Expert", 4), c("Low Coupling", 4),
                           c("High Cohesion", 4), c("GRASP", 4), c("Domain Model", 4),
                           c("Sequence Diagram", 7), c("Design Class Diagram", 4, 7)]}]  # fmt: skip
    assert practised_weeks(sheet) == [4]
    two_decks = [{"concepts": [c("A", 4), c("B", 4), c("C", 6), c("D", 6), c("E", 6)]}]
    assert practised_weeks(two_decks) == [4, 6]


def test_symbol_font_bullets_dont_reach_headings():
    from app.tutorials.blocks import _clean

    bullet = "\uf0b7"  # Word's Symbol-font bullet, a private-use character
    assert _clean(f"{bullet}  Why Does This Matter?") == "Why Does This Matter?"
