"""PPTX parsing: cached LibreOffice conversion and speaker notes. LibreOffice itself is
replaced by a stub that returns a generated PDF and counts conversions."""

import io
import uuid

import httpx
import pytest
from pptx import Presentation
from sqlalchemy import select

from app.db.models import Slide
from app.parse import pipeline as pipeline_module
from app.parse.pipeline import Parser
from app.services.files import delete_file
from app.services.uploads import ingest_upload
from app.storage.client import StorageError, StorageNotFound, _raise_for, converted_pdf_path
from tests.slide_fixtures import make_deck
from tests.test_pipeline import FakeVision


def make_pptx(slides: list[dict]) -> bytes:
    """Each dict: title, notes (optional), hidden (optional)."""
    prs = Presentation()
    for spec in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = spec["title"]
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
        if spec.get("hidden"):
            slide._element.set("show", "0")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def text_page(i: int) -> dict:
    return {"title": f"Topic {i}", "lines": [f"• Point {i}.{j}" for j in range(3)]}


# 6 slides, slide 3 hidden: the PDF export has 5 pages, and notes must skip the hidden one.
PPTX_SLIDES = [
    {"title": "Topic 0"},
    {"title": "Topic 1", "notes": "Say this about topic 1"},
    {"title": "Hidden draft", "notes": "never shown", "hidden": True},
    {"title": "Topic 2"},
    {"title": "Topic 3", "notes": "Low coupling is the point here"},
    {"title": "Topic 4"},
]
PDF_PAGES = [text_page(i) for i in range(5)]


@pytest.fixture
def converter(monkeypatch):
    calls = []
    pdf = make_deck(PDF_PAGES)

    def fake_convert(data: bytes) -> bytes:
        calls.append(len(data))
        return pdf

    monkeypatch.setattr(pipeline_module, "pptx_to_pdf", fake_convert)
    return calls


@pytest.fixture
def pptx_file(db, storage, subject):
    return ingest_upload(db, storage, subject, "Lecture 4 - GRASP.pptx", None,
                         make_pptx(PPTX_SLIDES), max_bytes=50_000_000).file  # fmt: skip


def parse(session_factory, storage, file_id):
    Parser(session_factory, storage, FakeVision).parse(file_id)


def rows(db, file_id) -> list[Slide]:
    db.expire_all()
    return list(db.scalars(select(Slide).where(Slide.file_id == file_id).order_by(Slide.page_no)))


def test_notes_are_stored_on_the_right_pages(db, session_factory, storage, pptx_file, converter):
    parse(session_factory, storage, pptx_file.id)
    notes = {r.page_no: r.notes for r in rows(db, pptx_file.id)}
    # Page 2 is slide 2; page 4 is slide 5 (slide 3 is hidden and not in the PDF).
    assert notes == {1: None, 2: "Say this about topic 1", 3: None,
                     4: "Low coupling is the point here", 5: None}  # fmt: skip
    # Notes live in their own column, not mixed into the parsed slide content.
    assert all("Low coupling" not in str(r.content_json) for r in rows(db, pptx_file.id))


def test_conversion_is_cached_by_pptx_hash(db, session_factory, storage, pptx_file, converter):
    parse(session_factory, storage, pptx_file.id)
    assert len(converter) == 1
    assert converted_pdf_path(pptx_file.sha256) in storage.objects

    parse(session_factory, storage, pptx_file.id)
    assert len(converter) == 1  # second parse reused the stored PDF
    db.refresh(pptx_file)
    assert pptx_file.status == "parsed"


def test_converted_pdf_is_deleted_with_the_last_copy(
    db, session_factory, storage, pptx_file, converter
):
    parse(session_factory, storage, pptx_file.id)
    delete_file(db, storage, pptx_file)
    assert converted_pdf_path(pptx_file.sha256) not in storage.objects


def test_notes_skipped_when_pages_and_slides_disagree(
    db, session_factory, storage, subject, monkeypatch
):
    monkeypatch.setattr(pipeline_module, "pptx_to_pdf", lambda _: make_deck(PDF_PAGES[:3]))
    file = ingest_upload(db, storage, subject, "Lecture 5.pptx", None, make_pptx(PPTX_SLIDES),
                         max_bytes=50_000_000).file  # fmt: skip
    parse(session_factory, storage, file.id)
    assert all(r.notes is None for r in rows(db, file.id))  # better none than misaligned


def test_viewer_and_api_show_notes(client, session_factory, storage, converter):
    r = client.post("/subjects", data={"name": "SMD", "code": "SWEN30006", "weeks": "12"},
                    follow_redirects=False)  # fmt: skip
    sid = r.headers["location"].rsplit("/", 1)[1]
    fid = client.post(
        f"/api/subjects/{sid}/files", files={"file": ("Lecture 4.pptx", make_pptx(PPTX_SLIDES))}
    ).json()["file"]["id"]
    parse(session_factory, storage, uuid.UUID(fid))

    page = client.get(f"/files/{fid}/slides?page=4").text
    assert "Speaker notes" in page and "Low coupling is the point here" in page
    assert "Speaker notes" not in client.get(f"/files/{fid}/slides?page=1").text
    api = client.get(f"/api/files/{fid}/slides").json()
    assert api[3]["notes"] == "Low coupling is the point here"


@pytest.mark.parametrize(
    ("status", "body", "error"),
    [
        (400, {"statusCode": "404", "error": "not_found", "message": "Object not found"},
         StorageNotFound),
        (404, {"message": "Not found"}, StorageNotFound),
        (400, {"statusCode": "400", "message": "Invalid key"}, StorageError),
        (500, {"message": "boom"}, StorageError),
    ],
)  # fmt: skip
def test_missing_objects_are_told_apart_from_errors(status, body, error):
    with pytest.raises(error) as info:
        _raise_for(httpx.Response(status, json=body), "download x")
    assert type(info.value) is error


def test_speaker_notes_are_not_sent_to_the_model(
    session_factory, storage, pptx_file, converter, monkeypatch
):
    # Make every page a vision page so each one's request can be inspected.
    monkeypatch.setattr(
        pipeline_module, "route_page", lambda page: pipeline_module.Decision("vision", "test")
    )
    vision = FakeVision()
    Parser(session_factory, storage, lambda: vision).parse(pptx_file.id)
    assert len(vision.calls) == 5
    sent = " ".join(str(call) for call in vision.calls)
    assert "Say this about topic 1" not in sent and "Low coupling" not in sent
