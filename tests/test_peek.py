import io

import pymupdf
from docx import Document
from pptx import Presentation

from app.classify.peek import peek


def _pdf(text: str, width: float, height: float) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=width, height=height)
    page.insert_text((50, 72), text)
    return doc.tobytes()


def test_pdf_text_and_orientation():
    landscape = peek(_pdf("Week 3: Decision Trees", 960, 540), "pdf")
    assert "Decision Trees" in landscape.text
    assert landscape.landscape is True
    assert peek(_pdf("Tutorial 2", 595, 842), "pdf").landscape is False


def test_pptx_first_slide_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Lecture 5: Naive Bayes"
    buf = io.BytesIO()
    prs.save(buf)
    result = peek(buf.getvalue(), "pptx")
    assert "Naive Bayes" in result.text
    assert result.landscape is True


def test_docx_paragraphs():
    doc = Document()
    doc.add_paragraph("Workshop 4")
    doc.add_paragraph("Question 1. Compute the entropy.")
    buf = io.BytesIO()
    doc.save(buf)
    assert "Question 1" in peek(buf.getvalue(), "docx").text


def test_text_lines_and_bom():
    result = peek("﻿WEBVTT\n\n00:00:01.000 --> 00:00:02.000\nHi".encode(), "vtt")
    assert result.lines[0] == "WEBVTT"


def test_corrupt_file_reports_error_instead_of_raising():
    result = peek(b"not a pdf", "pdf")
    assert result.text == ""
    assert result.error and "could not read" in result.error


def test_pdf_title_candidates_largest_text_first():
    doc = pymupdf.open()
    page = doc.new_page(width=960, height=540)
    page.insert_text((50, 60), "COMP30027 Machine Learning", fontsize=14)
    page.insert_text((50, 200), "Decision Trees", fontsize=36)
    page.insert_text((50, 300), "Dr Jane Smith", fontsize=12)
    result = peek(doc.tobytes(), "pdf")
    assert result.title_candidates[0] == "Decision Trees"


def test_pptx_title_candidates():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Naive Bayes"
    buf = io.BytesIO()
    prs.save(buf)
    assert peek(buf.getvalue(), "pptx").title_candidates == ("Naive Bayes",)


def test_docx_heading_is_first_candidate():
    doc = Document()
    doc.add_paragraph("COMP30027")
    doc.add_heading("Workshop 4: Entropy", level=1)
    buf = io.BytesIO()
    doc.save(buf)
    assert peek(buf.getvalue(), "docx").title_candidates[0] == "Workshop 4: Entropy"


def test_page_in_one_font_size_has_no_title():
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 72), "Examination\nReading time: 15 minutes\nWriting time: 3 hours")
    assert peek(doc.tobytes(), "pdf").title_candidates == ()


def test_sparse_title_slide_where_the_title_is_the_longest_text():
    doc = pymupdf.open()
    page = doc.new_page(width=960, height=540)
    page.insert_text((50, 60), "COMP30027 Machine Learning", fontsize=14)
    page.insert_text((50, 220), "Lecture 2: Probability Review", fontsize=34)
    assert peek(doc.tobytes(), "pdf").title_candidates[0] == "Lecture 2: Probability Review"
